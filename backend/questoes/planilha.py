"""
Importação de questões por planilha Excel (.xlsx) e geração da planilha modelo.

A planilha tem uma linha por questão. Os cabeçalhos são reconhecidos pelo nome
(maiúsculas, acentos e o asterisco não importam), então a ordem das colunas pode mudar.

O resultado de ler_planilha() tem o mesmo formato da importação de PDF,
e é gravado pela mesma função importacao.salvar_questoes().
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

# coluna interna -> (cabeçalho exibido, largura, obrigatória, comentário)
COLUNAS = [
    ("materia", "Matéria*", 22, True, "Ex.: Português, Direito Administrativo. Se não existir, é criada."),
    ("assunto", "Assunto", 24, False, "Ex.: Crase, Atos administrativos. Opcional."),
    ("tipo", "Tipo*", 17, True, "Múltipla escolha ou Certo/Errado."),
    ("enunciado", "Enunciado*", 60, True, "Texto da questão. Use Alt+Enter para quebrar linha dentro da célula."),
    ("a", "Alternativa A", 30, False, "Múltipla escolha: preencha de A em diante (mínimo 2). Certo/Errado: deixe vazio."),
    ("b", "Alternativa B", 30, False, None),
    ("c", "Alternativa C", 30, False, None),
    ("d", "Alternativa D", 30, False, None),
    ("e", "Alternativa E", 30, False, None),
    ("gabarito", "Gabarito*", 11, True, "Múltipla escolha: A, B, C, D ou E. Certo/Errado: Certo ou Errado."),
    ("explicacao", "Explicação", 45, False, "Aparece depois que você responde. Opcional, mas ajuda muito a fixar."),
    ("dificuldade", "Dificuldade", 13, False, "Fácil, Média ou Difícil (vale 10, 15 ou 20 XP). Vazio = Média."),
    ("banca", "Banca", 18, False, None),
    ("ano", "Ano", 8, False, None),
    ("orgao", "Órgão", 22, False, None),
    ("cargo", "Cargo", 22, False, None),
    ("codigo", "Código", 13, False, "Opcional. Código da questão na origem (ex.: Q2178820). Evita importar repetida."),
]

TIPOS = {"ME": "Múltipla escolha", "CE": "Certo/Errado"}
DIFICULDADES = {1: "Fácil", 2: "Média", 3: "Difícil"}
LETRAS = ["A", "B", "C", "D", "E"]

EXEMPLOS = [
    {
        "materia": "Direito Constitucional", "assunto": "Princípios Fundamentais", "tipo": "ME",
        "enunciado": "É um dos objetivos fundamentais da República Federativa do Brasil previsto "
                     "na Constituição Federal de 1988",
        "a": "a garantia do desenvolvimento nacional.", "b": "a soberania.",
        "c": "a dignidade da pessoa humana.", "d": "a cidadania.", "e": "o pluralismo político.",
        "gabarito": "A",
        "explicacao": "Objetivos fundamentais estão no art. 3º (ex.: garantir o desenvolvimento nacional). "
                      "Soberania, cidadania, dignidade e pluralismo são FUNDAMENTOS (art. 1º).",
        "dificuldade": "Fácil", "banca": "CESPE/CEBRASPE", "ano": 2023,
        "orgao": "AGER - Mato Grosso", "cargo": "Técnico Administrativo", "codigo": "Q2178841",
    },
    {
        "materia": "Administração Geral", "assunto": "Gestão de Pessoas", "tipo": "ME",
        "enunciado": "Nas organizações, é atribuição da gestão de pessoas\n"
                     "I garantir certo equilíbrio entre os interesses dos colaboradores e os da organização.\n"
                     "II desenvolver e manter qualidade de vida no trabalho.\n"
                     "III administrar a mudança.\n"
                     "IV avaliar candidatos e propostas de fornecedores de materiais, aplicando instrumentos "
                     "como provas e avaliações de perfil.\n"
                     "V viabilizar a integração de novos colaboradores, promovendo estímulos para motivá-los.\n"
                     "Estão certos apenas os itens",
        "a": "I e II.", "b": "III e IV.", "c": "I, IV e V.", "d": "I, II, III e V.", "e": "II, III, IV e V.",
        "gabarito": "D", "explicacao": "", "dificuldade": "Média", "banca": "CESPE/CEBRASPE",
        "ano": 2023, "orgao": "AGER - Mato Grosso", "cargo": "Técnico Administrativo", "codigo": "Q2179422",
    },
    {
        "materia": "Direito Administrativo", "assunto": "Princípios", "tipo": "CE",
        "enunciado": "O princípio da impessoalidade impede que a publicidade de atos e programas públicos "
                     "contenha nomes ou imagens que caracterizem promoção pessoal de autoridades.",
        "gabarito": "Certo",
        "explicacao": "Art. 37, §1º, da CF: a publicidade deve ter caráter educativo, informativo ou de "
                      "orientação social, sem promoção pessoal.",
        "dificuldade": "Fácil", "banca": "", "ano": None, "orgao": "", "cargo": "", "codigo": "",
    },
    {
        "materia": "Português", "assunto": "Concordância", "tipo": "CE",
        "enunciado": "Na frase \"Fazem dois anos que estudo para concursos\", a concordância verbal está correta.",
        "gabarito": "Errado",
        "explicacao": "\"Fazer\" indicando tempo decorrido é impessoal e fica no singular: \"Faz dois anos\".",
        "dificuldade": "Média", "banca": "", "ano": None, "orgao": "", "cargo": "", "codigo": "",
    },
]

MATERIAS_SUGERIDAS = [
    "Português", "Raciocínio Lógico", "Matemática", "Informática", "Direito Constitucional",
    "Direito Administrativo", "Direito Penal", "Direito Processual Penal", "Direito Civil",
    "Direito Tributário", "Administração Geral", "Administração Pública", "Arquivologia",
    "Contabilidade", "AFO", "Ética no Serviço Público", "Legislação Estadual", "Legislação Federal",
    "Atualidades", "Inglês",
]


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------
def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", texto.lower())


APELIDOS = {
    "materia": "materia", "disciplina": "materia", "assunto": "assunto", "tema": "assunto",
    "tipo": "tipo", "enunciado": "enunciado", "pergunta": "enunciado", "questao": "enunciado",
    "alternativaa": "a", "a": "a", "alternativab": "b", "b": "b", "alternativac": "c", "c": "c",
    "alternativad": "d", "d": "d", "alternativae": "e", "e": "e",
    "gabarito": "gabarito", "resposta": "gabarito", "correta": "gabarito",
    "explicacao": "explicacao", "comentario": "explicacao",
    "dificuldade": "dificuldade", "banca": "banca", "ano": "ano", "orgao": "orgao",
    "cargo": "cargo", "codigo": "codigo", "id": "codigo",
}


def _texto_celula(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).replace("\r\n", "\n").strip()


def _tipo(valor: str, alternativas: list[str], gabarito: str) -> str:
    v = _normalizar(valor)
    if v in ("ce", "certoerrado", "certoouerrado", "ceerto", "certo"):
        return "CE"
    if v.startswith("mult") or v == "me":
        return "ME"
    # vazio: deduz pelo conteúdo
    if not any(alternativas) and _normalizar(gabarito) in ("c", "e", "certo", "errado"):
        return "CE"
    return "ME"


def _dificuldade(valor: str) -> int:
    v = _normalizar(valor)
    if v in ("1", "facil"):
        return 1
    if v in ("3", "dificil"):
        return 3
    return 2


def _encontrar_aba_e_cabecalho(wb):
    for ws in wb.worksheets:
        for linha in ws.iter_rows(min_row=1, max_row=10):
            mapa = {}
            for celula in linha:
                chave = APELIDOS.get(_normalizar(celula.value))
                if chave and chave not in mapa:
                    mapa[chave] = celula.column - 1
            if "enunciado" in mapa and "gabarito" in mapa:
                return ws, linha[0].row, mapa
    return None, None, None


def ler_planilha(caminho) -> dict:
    from openpyxl import load_workbook

    from .importacao import ImportacaoError

    try:
        wb = load_workbook(caminho, data_only=True, read_only=True)
    except Exception as exc:
        raise ImportacaoError(f"Não consegui abrir a planilha: {exc}")

    ws, linha_cab, mapa = _encontrar_aba_e_cabecalho(wb)
    if ws is None:
        raise ImportacaoError(
            "Não encontrei os cabeçalhos da planilha. Ela precisa ter pelo menos as colunas "
            "'Enunciado' e 'Gabarito' (use a planilha modelo)."
        )

    questoes = []
    for valores in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
        def col(nome):
            i = mapa.get(nome)
            return _texto_celula(valores[i]) if i is not None and i < len(valores) else ""

        if not any(col(c) for c in ("materia", "enunciado", "a", "gabarito")):
            continue  # linha vazia
        numero_linha = linha_cab + 1 + len(questoes)
        avisos = []

        textos_alt = [col(l.lower()) for l in LETRAS]
        gabarito_bruto = col("gabarito")
        tipo = _tipo(col("tipo"), textos_alt, gabarito_bruto)

        if tipo == "CE":
            alternativas = [{"letra": "Certo", "texto": "Certo"}, {"letra": "Errado", "texto": "Errado"}]
            g = _normalizar(gabarito_bruto)
            gabarito = {"c": "Certo", "certo": "Certo", "e": "Errado", "errado": "Errado"}.get(g)
            if any(textos_alt):
                avisos.append("Questão Certo/Errado com alternativas preenchidas: as alternativas foram ignoradas.")
        else:
            ultima = max((i for i, t in enumerate(textos_alt) if t), default=-1)
            alternativas = [{"letra": LETRAS[i], "texto": textos_alt[i]} for i in range(ultima + 1)]
            if any(not a["texto"] for a in alternativas):
                avisos.append("Há alternativa vazia no meio (ex.: A e C preenchidas, B vazia).")
            g = gabarito_bruto.strip().upper()[:1]
            gabarito = g if g in LETRAS else None
            if len(alternativas) < 2:
                avisos.append("Múltipla escolha precisa de pelo menos 2 alternativas.")

        if not gabarito:
            avisos.append(f"Gabarito inválido ou vazio: '{gabarito_bruto}'.")
        elif gabarito not in [a["letra"] for a in alternativas]:
            avisos.append(f"Gabarito '{gabarito}' aponta para alternativa vazia.")
        if not col("materia"):
            avisos.append("Matéria não preenchida.")
        if not col("enunciado"):
            avisos.append("Enunciado vazio.")

        ano = col("ano")
        questoes.append({
            "numero": numero_linha,
            "codigo_externo": col("codigo"),
            "materia": col("materia") or None,
            "assunto": col("assunto") or None,
            "banca": col("banca"),
            "ano": int(ano) if ano.isdigit() else None,
            "orgao": col("orgao"),
            "cargo": col("cargo"),
            "prova": "",
            "tipo": tipo,
            "enunciado": col("enunciado"),
            "alternativas": alternativas,
            "gabarito": gabarito,
            "explicacao": col("explicacao"),
            "dificuldade": _dificuldade(col("dificuldade")),
            "avisos": [f"Linha {numero_linha}: {a}" for a in avisos],
        })
    wb.close()

    avisos_gerais = [] if questoes else ["A planilha não tem nenhuma questão preenchida."]
    return {"questoes": questoes, "avisos": avisos_gerais, "usou_ocr": False, "origem": "planilha"}


# ---------------------------------------------------------------------------
# Planilha modelo
# ---------------------------------------------------------------------------
def gerar_modelo(destino, com_exemplos: bool = True):
    """Gera a planilha modelo (caminho ou arquivo aberto em modo binário)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fonte = "Arial"
    borda = Border(*(Side(style="thin", color="D5DAE5"),) * 4)
    wb = Workbook()

    # --- Aba Questões ---
    ws = wb.active
    ws.title = "Questões"
    for i, (_, titulo, largura, obrigatoria, comentario) in enumerate(COLUNAS, start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = Font(name=fonte, bold=True, color="FFFFFF" if obrigatoria else "1B1F2E", size=10)
        c.fill = PatternFill("solid", fgColor="A6670A" if obrigatoria else "F5D58A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
        if comentario:
            c.comment = Comment(comentario, "Modelo")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    if com_exemplos:
        for r, ex in enumerate(EXEMPLOS, start=2):
            for i, (chave, *_resto) in enumerate(COLUNAS, start=1):
                valor = ex.get(chave)
                if chave == "tipo":
                    valor = TIPOS[valor]
                ws.cell(row=r, column=i, value=valor if valor not in ("", None) else None)

    ultima_linha = 1000
    for r in range(2, ultima_linha + 1):
        for i in range(1, len(COLUNAS) + 1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=fonte, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if r <= 60:
                c.border = borda
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ultima_linha}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # --- Aba Listas (fonte das listas suspensas) ---
    listas = wb.create_sheet("Listas")
    cabecalhos = ["Tipos", "Gabaritos", "Dificuldades", "Matérias sugeridas"]
    valores = [list(TIPOS.values()), LETRAS + ["Certo", "Errado"], list(DIFICULDADES.values()),
               MATERIAS_SUGERIDAS]
    for j, (cab, vals) in enumerate(zip(cabecalhos, valores), start=1):
        listas.cell(row=1, column=j, value=cab).font = Font(name=fonte, bold=True)
        for k, v in enumerate(vals, start=2):
            listas.cell(row=k, column=j, value=v).font = Font(name=fonte)
        listas.column_dimensions[get_column_letter(j)].width = 26

    def coluna(chave):
        return get_column_letter([c[0] for c in COLUNAS].index(chave) + 1)

    def validar(chave, ref, estrita=True, msg=""):
        dv = DataValidation(type="list", formula1=ref, allow_blank=True)
        dv.showErrorMessage = estrita
        dv.error = msg
        dv.errorTitle = "Valor inválido"
        ws.add_data_validation(dv)
        dv.add(f"{coluna(chave)}2:{coluna(chave)}{ultima_linha}")

    validar("tipo", f"=Listas!$A$2:$A${1 + len(TIPOS)}", msg="Escolha Múltipla escolha ou Certo/Errado.")
    validar("gabarito", f"=Listas!$B$2:$B$8", msg="Use A, B, C, D, E, Certo ou Errado.")
    validar("dificuldade", f"=Listas!$C$2:$C$4", msg="Escolha Fácil, Média ou Difícil.")
    # matéria: sugere a lista, mas aceita qualquer nome
    validar("materia", f"=Listas!$D$2:$D${1 + len(MATERIAS_SUGERIDAS)}", estrita=False)

    # --- Aba Como preencher ---
    ajuda = wb.create_sheet("Como preencher", 0)
    ajuda.column_dimensions["A"].width = 3
    ajuda.column_dimensions["B"].width = 26
    ajuda.column_dimensions["C"].width = 90
    linhas = [
        ("titulo", "Planilha de questões — Estudo Gamificado"),
        ("texto", "Preencha a aba Questões: uma linha por questão. Depois importe em "
                  "Admin › Questões › Importar prova."),
        ("vazio", ""),
        ("secao", "Legenda dos cabeçalhos"),
        ("cor_obrig", "Cabeçalho dourado escuro = coluna obrigatória (marcada com *)."),
        ("cor_opc", "Cabeçalho dourado claro = coluna opcional."),
        ("texto", "Passe o mouse sobre cada cabeçalho para ver a dica da coluna."),
        ("vazio", ""),
        ("secao", "Regras"),
    ]
    for chave, titulo, _l, obrig, coment in COLUNAS:
        if coment:
            linhas.append(("regra", ("Alternativas A a E" if chave == "a" else titulo, coment)))
    linhas += [
        ("vazio", ""),
        ("secao", "Dicas"),
        ("texto", "As 4 linhas da aba Questões são exemplos no formato certo: apague-as ou substitua pelas suas."),
        ("texto", "Questões repetidas não duplicam: o app compara o Código ou, se vazio, o enunciado."),
        ("texto", "Para entrar direto no jogo, marque 'Publicar automaticamente' na tela de importação. "
                  "Linhas com problema ficam em rascunho, com o aviso indicando a linha."),
        ("texto", "Não renomeie os cabeçalhos da linha 1. A ordem das colunas pode mudar."),
    ]
    r = 2
    for tipo, conteudo in linhas:
        if tipo in ("titulo", "texto"):
            ajuda.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        if tipo == "titulo":
            c = ajuda.cell(row=r, column=2, value=conteudo)
            c.font = Font(name=fonte, bold=True, size=16, color="A6670A")
        elif tipo == "secao":
            c = ajuda.cell(row=r, column=2, value=conteudo)
            c.font = Font(name=fonte, bold=True, size=12)
        elif tipo == "regra":
            ajuda.cell(row=r, column=2, value=conteudo[0]).font = Font(name=fonte, bold=True, size=10)
            c = ajuda.cell(row=r, column=3, value=conteudo[1])
            c.font = Font(name=fonte, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        elif tipo in ("cor_obrig", "cor_opc"):
            amostra = ajuda.cell(row=r, column=2, value="Matéria*" if tipo == "cor_obrig" else "Assunto")
            amostra.fill = PatternFill("solid", fgColor="A6670A" if tipo == "cor_obrig" else "F5D58A")
            amostra.font = Font(name=fonte, bold=True, size=10,
                                color="FFFFFF" if tipo == "cor_obrig" else "1B1F2E")
            ajuda.cell(row=r, column=3, value=conteudo).font = Font(name=fonte, size=10)
        elif tipo == "texto":
            c = ajuda.cell(row=r, column=2, value=conteudo)
            c.font = Font(name=fonte, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if len(conteudo) > 110:
                ajuda.row_dimensions[r].height = 27
        r += 1
    ajuda.sheet_view.showGridLines = False
    ajuda.page_setup.fitToWidth = 1
    ajuda.page_setup.fitToHeight = 0
    ajuda.sheet_properties.pageSetUpPr.fitToPage = True

    wb.active = wb.sheetnames.index("Questões")
    wb.save(destino)
    return destino


def gerar_modelo_em_disco(pasta) -> Path:
    caminho = Path(pasta) / "modelo_questoes.xlsx"
    gerar_modelo(caminho)
    return caminho
