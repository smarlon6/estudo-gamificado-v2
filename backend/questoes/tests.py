"""Testes do importador com dados sintéticos (não dependem de PDF nem de OCR)."""
from django.test import TestCase

from .importacao import (
    Circulo,
    Linha,
    Palavra,
    ResultadoLeitura,
    _corrigir_itens_romanos,
    _corrigir_ocr,
    _interpretar_gabarito,
    _parse_classificacao,
    interpretar,
    salvar_questoes,
)
from .models import Questao


def linha(pagina, top, texto, x0=34.0):
    palavras, x = [], x0
    for p in texto.split():
        palavras.append(Palavra(p, x, x + 5 * len(p), top, top + 9))
        x += 5 * len(p) + 3
    return Linha(pagina, palavras)


def prova_sintetica(com_gabarito=True):
    linhas = [
        linha(0, 100, "1 Q1234567 Português > Crase, Uso da crase"),
        linha(0, 115, "Ano: 2023 Banca: CESPE / CEBRASPE Órgão: AGER - Mato Grosso"),
        linha(0, 128, "Prova: CESPE / CEBRASPE - 2023 - AGER - Mato Grosso - Técnico Administrativo"),
        linha(0, 150, "Assinale a frase correta."),
        linha(0, 172, "(A) Vou à escola.", x0=36), linha(0, 190, "Vou a escola.", x0=48),
        linha(0, 208, "Vou á escola.", x0=48), linha(0, 226, "Vou à pé.", x0=48),
        linha(0, 244, "Vou àquela.", x0=48),
        linha(1, 60, "2 02234567 Não definido"),          # OCR leu o Q como 0
        linha(1, 75, "Ano: 2023 Banca: CESPE / CEBRASPE Órgão: AGER - Mato Grosso"),
        linha(1, 88, "Prova: CESPE / CEBRASPE - 2023 - AGER - Mato Grosso - Técnico Administrativo"),
        linha(1, 110, "Julgue o item."),
        linha(1, 132, "Certo", x0=48), linha(1, 150, "Errado", x0=48),
    ]
    if com_gabarito:
        linhas += [linha(1, 300, "Respostas"), linha(1, 318, "1D 2:C")]  # "1D": OCR comeu o ":"
    circulos = [Circulo(0, 170 + 18 * i, 182 + 18 * i, 46) for i in range(5)]
    circulos += [Circulo(1, 130, 142, 46), Circulo(1, 148, 160, 46)]
    return ResultadoLeitura(linhas, circulos, usou_ocr=True, paginas=2)


class InterpretacaoTest(TestCase):
    def test_interpreta_questoes_e_gabarito(self):
        dados = interpretar(prova_sintetica())
        q1, q2 = dados["questoes"]

        self.assertEqual(q1["codigo_externo"], "Q1234567")
        self.assertEqual(q1["materia"], "Português")
        self.assertEqual(q1["assunto"], "Crase")
        self.assertEqual(q1["ano"], 2023)
        self.assertEqual(q1["cargo"], "Técnico Administrativo")
        self.assertEqual(q1["enunciado"], "Assinale a frase correta.")
        self.assertEqual([a["letra"] for a in q1["alternativas"]], list("ABCDE"))
        self.assertEqual(q1["alternativas"][0]["texto"], "Vou à escola.")  # "(A)" removido
        self.assertEqual(q1["gabarito"], "D")

        self.assertEqual(q2["codigo_externo"], "Q2234567")
        self.assertEqual(q2["tipo"], "CE")
        self.assertEqual(q2["gabarito"], "Certo")
        self.assertIn("Matéria não definida", " ".join(q2["avisos"]))

    def test_sem_gabarito_gera_aviso(self):
        dados = interpretar(prova_sintetica(com_gabarito=False))
        self.assertIn("Gabarito não encontrado.", dados["questoes"][0]["avisos"])

    def test_gabarito_com_erros_de_ocr(self):
        self.assertEqual(_interpretar_gabarito("1D 2: E 3:A 40:C 5:D"),
                         {1: "D", 2: "E", 3: "A", 4: "C", 5: "D"})

    def test_classificacao(self):
        self.assertEqual(_parse_classificacao("Direito Administrativo > Atos administrativos, Atributos"),
                         ("Direito Administrativo", "Atos administrativos"))
        self.assertEqual(_parse_classificacao("Não definido"), (None, None))

    def test_correcoes_de_romanos(self):
        self.assertEqual(_corrigir_ocr("apenas os itens |, Il e Ill."), "apenas os itens I, II e III.")
        textos, inicios = _corrigir_itens_romanos(
            ["Julgue:", "| primeiro item", "Il segundo", "continua", "INI terceiro", "IV quarto"]
        )
        self.assertEqual(textos[1:], ["I primeiro item", "II segundo", "continua",
                                      "III terceiro", "IV quarto"])
        self.assertEqual(inicios, {1, 2, 4, 5})


class GravacaoTest(TestCase):
    def test_salva_como_rascunho_e_nao_duplica(self):
        dados = interpretar(prova_sintetica())
        r1 = salvar_questoes(dados)
        self.assertEqual(r1["criadas"], 2)
        self.assertEqual(Questao.objects.filter(status="rascunho").count(), 2)
        q = Questao.objects.get(codigo_externo="Q1234567")
        self.assertEqual(q.alternativa_correta.letra, "D")
        self.assertEqual(q.materia.nome, "Português")

        r2 = salvar_questoes(dados)
        self.assertEqual(r2["duplicadas"], 2)
        self.assertEqual(Questao.objects.count(), 2)

    def test_publicar_so_ativa_quem_nao_tem_aviso(self):
        salvar_questoes(interpretar(prova_sintetica()), publicar=True)
        self.assertEqual(Questao.objects.get(codigo_externo="Q1234567").status, "ativa")
        sem_materia = Questao.objects.get(codigo_externo="Q2234567")
        self.assertEqual(sem_materia.status, "rascunho")
        self.assertEqual(sem_materia.materia.nome, "Não classificada")


class PlanilhaTest(TestCase):
    def setUp(self):
        import tempfile
        from pathlib import Path

        from openpyxl import load_workbook

        from .planilha import gerar_modelo

        self.pasta = Path(tempfile.mkdtemp())
        self.caminho = self.pasta / "modelo.xlsx"
        gerar_modelo(self.caminho)
        self.load_workbook = load_workbook

    def adicionar_linha(self, valores: dict):
        """Escreve uma linha nova na aba Questões usando os nomes das colunas internas."""
        from .planilha import COLUNAS

        wb = self.load_workbook(self.caminho)
        ws = wb["Questões"]
        linha = 2
        while ws.cell(row=linha, column=4).value:
            linha += 1
        for i, (chave, *_r) in enumerate(COLUNAS, start=1):
            if chave in valores:
                ws.cell(row=linha, column=i, value=valores[chave])
        wb.save(self.caminho)
        return linha

    def test_le_os_exemplos_do_modelo(self):
        from .importacao import ler_prova

        dados = ler_prova(self.caminho)
        q = dados["questoes"]
        self.assertEqual(len(q), 4)
        self.assertTrue(all(not x["avisos"] for x in q), [x["avisos"] for x in q])
        self.assertEqual(q[0]["gabarito"], "A")
        self.assertEqual(q[0]["dificuldade"], 1)
        self.assertEqual(len(q[0]["alternativas"]), 5)
        self.assertIn("\nII desenvolver", q[1]["enunciado"])  # quebra de linha preservada
        self.assertEqual(q[2]["tipo"], "CE")
        self.assertEqual(q[3]["gabarito"], "Errado")

    def test_valida_linhas_com_problema(self):
        from .importacao import ler_prova

        linha = self.adicionar_linha({
            "materia": "Informática", "tipo": "Múltipla escolha", "enunciado": "Pergunta?",
            "a": "um", "b": "dois", "gabarito": "D",
        })
        self.adicionar_linha({"enunciado": "Sem matéria", "gabarito": "C"})  # tipo vazio → Certo/Errado
        q = ler_prova(self.caminho)["questoes"]
        self.assertIn(f"Linha {linha}: Gabarito 'D' aponta para alternativa vazia.", q[4]["avisos"])
        self.assertEqual(q[5]["tipo"], "CE")
        self.assertEqual(q[5]["gabarito"], "Certo")
        self.assertTrue(any("Matéria não preenchida" in a for a in q[5]["avisos"]))

    def test_importa_publica_e_nao_duplica(self):
        from .importacao import ler_prova

        dados = ler_prova(self.caminho)
        r1 = salvar_questoes(dados, publicar=True)
        self.assertEqual((r1["criadas"], r1["ativas"]), (4, 4))
        q = Questao.objects.get(codigo_externo="Q2178841")
        self.assertEqual(q.dificuldade, 1)
        self.assertIn("art. 3º", q.explicacao)
        self.assertEqual(q.alternativa_correta.texto, "a garantia do desenvolvimento nacional.")
        # reimportar a mesma planilha não duplica (por código ou por enunciado)
        r2 = salvar_questoes(ler_prova(self.caminho), publicar=True)
        self.assertEqual((r2["criadas"], r2["duplicadas"]), (0, 4))
